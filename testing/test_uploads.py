"""Uploaded files: parsing, type inference, keys and links, preview edits (no database)."""

import io
import sys
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.uploads import (
    UploadError, analyse, analyse_column, apply_edits, clean_name, convert, parse_csv, parse_date,
    parse_files, parse_timestamp, to_preview,
)


def _csv(text, name="data.csv", encoding="utf-8"):
    return parse_csv(name, text.encode(encoding))


# --------------------------------------------------------------------------- #
# Names
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("raw, expected", [
    ("Order Date", "order_date"), ("orderDate", "order_date"), ("  Total (USD) ", "total_usd"),
    ("Été 2025", "ete_2025"), ("1st place", "c_1st_place"), ("", "fallback"), ("***", "fallback"),
    ("x" * 80, "x" * 63),
])
def test_names_are_cleaned_into_readable_identifiers(raw, expected):
    assert clean_name(raw, "fallback") == expected


def test_table_and_column_names_are_unique():
    tables = parse_files([("Sales.csv", b"a,A,a \n1,2,3\n"), ("sales.csv", b"x\n1\n")])
    assert [t.name for t in tables] == ["sales", "sales_2"]
    plan = analyse(tables)
    assert [c.name for c in plan.tables[0].columns] == ["a", "a_2", "a_3"]


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #

def test_semicolons_tabs_and_byte_order_marks_are_handled():
    semi = _csv("id;name\n1;Ada\n2;Grace\n")
    assert semi.headers == ["id", "name"] and semi.rows == [("1", "Ada"), ("2", "Grace")]
    assert "semicolons" in semi.notes[0]
    tabs = _csv("id\tname\n1\tAda\n")
    assert tabs.rows == [("1", "Ada")]
    bom = parse_csv("d.csv", b"\xef\xbb\xbfid,name\n1,Ada\n")
    assert bom.headers == ["id", "name"], "the BOM isn't part of the first header"


def test_windows_1252_files_are_read_and_flagged():
    table = _csv("name\nJosé\n", encoding="cp1252")
    assert table.rows == [("José",)]
    assert "Windows-1252" in table.notes[0]


def test_ragged_rows_are_padded_or_left_out_and_reported():
    table = _csv("a,b,c\n1,2\n1,2,3,4\n1,2,3,,\n\n5,6,7\n")
    assert table.rows == [("1", "2", None), ("1", "2", "3"), ("5", "6", "7")]
    assert any("fewer values" in n for n in table.notes)
    assert any("left out" in n and "line 3" in n for n in table.notes)


def test_empty_cells_are_null_and_values_are_trimmed():
    assert _csv("a,b\n  x  ,\n").rows == [("x", None)]


def test_empty_and_header_only_files():
    with pytest.raises(UploadError, match="empty"):
        _csv("\n\n")
    assert any("no data" in n for n in _csv("a,b\n").notes)


def test_limits_are_enforced(monkeypatch):
    import app.uploads as uploads
    monkeypatch.setattr(uploads, "MAX_ROWS_PER_TABLE", 3)
    with pytest.raises(UploadError, match="more than 3 rows"):
        _csv("a\n1\n2\n3\n4\n")
    with pytest.raises(UploadError, match="at most 10 files"):
        parse_files([(f"f{i}.csv", b"a\n1\n") for i in range(11)])
    monkeypatch.setattr(uploads, "MAX_TOTAL_BYTES", 10)
    with pytest.raises(UploadError, match="limit is"):
        parse_files([("a.csv", b"a\n" + b"1\n" * 20)])


@pytest.mark.parametrize("name, message", [
    ("dump.sql", "SQL files can't be uploaded"), ("old.xls", "save it as .xlsx"), ("notes.pdf", "only CSV and Excel"),
])
def test_unsupported_files_say_what_to_do(name, message):
    with pytest.raises(UploadError, match=message):
        parse_files([(name, b"x")])


# --------------------------------------------------------------------------- #
# Types
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("values, kind", [
    (["1", "-2", "30", None], "integer"),
    (["007", "12"], "text"),                     # leading zeros: codes, not numbers
    (["1.5", "2", "-0.25", "1e3"], "decimal"),
    (["1,234.50", "999"], "decimal"),
    (["yes", "No", "Y", "n"], "boolean"),
    (["1", "0"], "integer"),                     # 1/0 stay numbers
    (["2025-04-03", "2025-12-31"], "date"),
    (["3 Apr 2025", "April 4, 2025", "5-Sept-2025"], "date"),
    (["2025-04-03 14:30", "2025-04-03T09:00:00"], "timestamp"),
    (["2025-04-03", "2025-04-03 10:00"], "timestamp"),
    (["abc", "1"], "text"),
    (["99999999999999999999"], "text"),         # past BIGINT: an identifier, not arithmetic
    ([None, None], "text"),
])
def test_the_narrowest_type_that_fits_every_value(values, kind):
    assert analyse_column(values).type == kind


def test_ambiguous_dates_are_day_first_and_flagged():
    info = analyse_column(["03/04/2025", "05/06/2025"])
    assert (info.type, info.date_order, info.ambiguous_dates) == ("date", "dmy", True)
    assert convert("03/04/2025", "date", info.date_order) == date(2025, 4, 3)
    assert convert("03/04/2025", "date", "mdy") == date(2025, 3, 4)


def test_one_unambiguous_date_decides_the_whole_column():
    day_first = analyse_column(["03/04/2025", "25/04/2025"])
    assert (day_first.date_order, day_first.ambiguous_dates) == ("dmy", False)
    month_first = analyse_column(["03/04/2025", "04/25/2025"])
    assert (month_first.date_order, month_first.ambiguous_dates) == ("mdy", False)
    assert analyse_column(["25/04/2025", "04/25/2025"]).type == "text", "both orders in one column: not dates"


def test_times_with_a_zone_become_utc():
    info = analyse_column(["2025-04-03T10:00:00+05:30"])
    assert info.type == "timestamp" and any("UTC" in n for n in info.notes)
    assert parse_timestamp("2025-04-03T10:00:00+05:30") == datetime(2025, 4, 3, 4, 30)
    assert parse_timestamp("2025-04-03T10:00:00Z") == datetime(2025, 4, 3, 10, 0)


def test_impossible_dates_are_not_dates():
    assert parse_date("2025-02-30") is None
    assert analyse_column(["2025-02-30"]).type == "text"


def test_native_excel_values():
    assert analyse_column([1, 2, None]).type == "integer"
    assert analyse_column([1.0, 2.0]).type == "integer"
    assert analyse_column([1.5, 2]).type == "decimal"
    assert analyse_column([True, False]).type == "boolean"
    assert analyse_column([datetime(2025, 1, 1), datetime(2025, 2, 1)]).type == "date"
    assert analyse_column([datetime(2025, 1, 1, 9, 30)]).type == "timestamp"


@pytest.mark.parametrize("value, kind, expected", [
    ("1,234.50", "decimal", Decimal("1234.50")), (3.0, "text", "3"), ("Yes", "boolean", True),
    (datetime(2025, 1, 2, 3, 4), "date", date(2025, 1, 2)), ("2025-01-02", "timestamp", datetime(2025, 1, 2)),
    (None, "integer", None),
])
def test_conversion(value, kind, expected):
    assert convert(value, kind) == expected


@pytest.mark.parametrize("value, kind", [("1.5", "integer"), ("abc", "decimal"), ("maybe", "boolean"), ("2025-13-01", "date"), (True, "integer")])
def test_conversion_refuses_values_that_dont_fit(value, kind):
    with pytest.raises(ValueError):
        convert(value, kind)


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #

def _workbook(sheets):
    import openpyxl
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_each_sheet_becomes_a_table_with_native_types():
    data = _workbook({
        "Customers": [["Customer ID", "Name", "Joined", "VIP"], [1, "Ada", datetime(2025, 1, 5), True], [2, "Grace", datetime(2025, 2, 6), False]],
        "Empty": [],
        "Orders": [["order_id", "customer_id", "amount"], [10, 1, 99.5], [11, 2, 15]],
    })
    tables = parse_files([("shop.xlsx", data)])
    assert [t.name for t in tables] == ["customers", "orders"], "empty sheets are skipped"
    plan = analyse(tables)
    types = {c.name: c.type for c in plan.tables[0].columns}
    assert types == {"customer_id": "integer", "name": "text", "joined": "date", "vip": "boolean"}
    assert {c.name: c.type for c in plan.tables[1].columns}["amount"] == "decimal"
    assert [(l.table, l.column, l.ref_table) for l in plan.links] == [("orders", "customer_id", "customers")]


def test_a_single_sheet_workbook_is_named_after_the_file():
    tables = parse_files([("Q3 Sales.xlsx", _workbook({"Sheet1": [["a"], [1]]}))])
    assert tables[0].name == "q3_sales" and tables[0].source == "Q3 Sales.xlsx › Sheet1"


def _zip(members):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def test_workbooks_with_xml_entities_are_refused():
    evil = _zip({"[Content_Types].xml": b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><x>&a;</x>'})
    with pytest.raises(UploadError, match="XML declarations"):
        parse_files([("evil.xlsx", evil)])


def test_zip_bombs_are_refused():
    bomb = _zip({"xl/worksheets/sheet1.xml": b"0" * (30 * 1024 * 1024)})  # 30 MB of zeros packs to ~30 KB
    with pytest.raises(UploadError, match="too large to read safely"):
        parse_files([("bomb.xlsx", bomb)])


# --------------------------------------------------------------------------- #
# Keys, links, edits
# --------------------------------------------------------------------------- #

DEPARTMENTS = "department_id,name\n1,Sales\n2,Research\n"
EMPLOYEES = "employee_id,name,department_id,hired\n10,Ada,1,03/04/2024\n11,Grace,2,05/06/2024\n12,Linus,,07/08/2024\n"


def _plan(*files):
    raws = parse_files(list(files))
    return analyse(raws), raws


def test_primary_keys_and_fully_matching_links_are_found():
    plan, _ = _plan(("departments.csv", DEPARTMENTS.encode()), ("employees.csv", EMPLOYEES.encode()))
    assert {t.name: t.primary_key for t in plan.tables} == {"departments": "department_id", "employees": "employee_id"}
    assert [(l.table, l.column, l.ref_table, l.ref_column) for l in plan.links] == [
        ("employees", "department_id", "departments", "department_id")
    ]


def test_a_link_with_unmatched_values_is_not_made_and_says_why():
    employees = EMPLOYEES.replace("12,Linus,,", "12,Linus,9,")
    plan, _ = _plan(("departments.csv", DEPARTMENTS.encode()), ("employees.csv", employees.encode()))
    assert plan.links == []
    assert any("Not linked: employees.department_id" in n and "1 value has no match" in n for n in plan.notes)


def test_bare_id_columns_never_link_to_each_other():
    plan, _ = _plan(("a.csv", b"id,x\n1,a\n2,b\n"), ("b.csv", b"id,y\n1,c\n"))
    assert {t.name: t.primary_key for t in plan.tables} == {"a": "id", "b": "id"}
    assert plan.links == []


def test_a_singular_id_column_links_to_a_bare_id_key():
    plan, _ = _plan(("teams.csv", b"id,name\n1,Red\n2,Blue\n"), ("players.csv", b"id,name,team_id\n7,Ann,1\n8,Bo,2\n"))
    assert [(l.table, l.column, l.ref_table, l.ref_column) for l in plan.links] == [("players", "team_id", "teams", "id")]


def test_edits_rename_retype_and_relink():
    plan, raws = _plan(("departments.csv", DEPARTMENTS.encode()), ("staff.csv", b"staff_no,dept,hired\n1,1,03/04/2024\n2,2,05/06/2024\n"))
    assert plan.links == []
    staff = next(t for t in plan.tables if t.name == "staff")
    dept_key = next(c.key for c in staff.columns if c.name == "dept")
    hired_key = next(c.key for c in staff.columns if c.name == "hired")
    apply_edits(plan, raws, {"tables": [{"key": staff.key, "name": "Employees", "columns": [
        {"key": dept_key, "name": "department_id"}, {"key": hired_key, "date_order": "mdy"},
    ]}]})
    assert [t.name for t in plan.tables] == ["departments", "employees"]
    assert [(l.table, l.column, l.ref_table) for l in plan.links] == [("employees", "department_id", "departments")]
    hired = next(c for c in plan.tables[1].columns if c.key == hired_key)
    assert convert(raws[1].rows[0][2], hired.type, hired.date_order) == date(2024, 3, 4)


def test_edits_that_dont_fit_the_data_are_all_reported_at_once():
    plan, raws = _plan(("t.csv", b"a,b\nx,1.5\ny,2\n"))
    columns = plan.tables[0].columns
    with pytest.raises(UploadError) as info:
        apply_edits(plan, raws, {"tables": [{"key": "t0", "columns": [
            {"key": columns[0].key, "type": "integer"}, {"key": columns[1].key, "type": "integer"},
        ]}]})
    message = str(info.value)
    assert "t.a: 2 values don't fit the type integer (for example 'x')" in message
    assert "t.b: 1 value doesn't fit the type integer (for example '1.5')" in message


def test_leaving_out_tables_and_duplicate_names():
    plan, raws = _plan(("a.csv", b"x\n1\n"), ("b.csv", b"y\n2\n"))
    with pytest.raises(UploadError, match="at least one table"):
        apply_edits(plan, raws, {"tables": [{"key": "t0", "include": False}, {"key": "t1", "include": False}]})
    plan, raws = _plan(("a.csv", b"x\n1\n"), ("b.csv", b"y\n2\n"))
    with pytest.raises(UploadError, match="Two tables are named a"):
        apply_edits(plan, raws, {"tables": [{"key": "t1", "name": "a"}]})


def test_the_preview_carries_everything_the_editor_needs():
    plan, _ = _plan(("departments.csv", DEPARTMENTS.encode()), ("employees.csv", EMPLOYEES.encode()))
    preview = to_preview(plan)
    employees = preview["tables"][1]
    hired = next(c for c in employees["columns"] if c["name"] == "hired")
    assert hired == {**hired, "type": "date", "date_order": "dmy", "ambiguous_dates": True, "original": "hired"}
    assert hired["samples"] == ["03/04/2024", "05/06/2024", "07/08/2024"]
    assert employees["rows"] == 3 and employees["primary_key"] == "employee_id"
    assert preview["links"][0]["ref_table"] == "departments" and "integer" in preview["types"]


def test_renaming_a_table_keeps_its_key_and_links():
    # Seen in the browser check: employees renamed to staff lost employee_id as key and its links.
    assignments = b"assignment_id,employee_id\n1,10\n2,11\n"
    plan, raws = _plan(("employees.csv", EMPLOYEES.encode()), ("assignments.csv", assignments))
    employees = next(t for t in plan.tables if t.name == "employees")
    apply_edits(plan, raws, {"tables": [{"key": employees.key, "name": "staff"}]})
    staff = next(t for t in plan.tables if t.name == "staff")
    assert staff.primary_key == "employee_id"
    assert ("assignments", "employee_id", "staff", "employee_id") in [(l.table, l.column, l.ref_table, l.ref_column) for l in plan.links]


def test_a_first_column_ending_in_id_is_the_key():
    plan, _ = _plan(("people.csv", b"person_code_id,name\nA1,Ann\nB2,Bo\n"))
    assert plan.tables[0].primary_key == "person_code_id"
    plan, _ = _plan(("people.csv", b"name,ref_id\nAnn,1\nBo,1\n"))
    assert plan.tables[0].primary_key is None, "only the first column, and only when unique"
