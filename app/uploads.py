"""
Uploaded files -> tables in a workspace.

    parse      CSV (encoding and delimiter detected) or XLSX (one table per
               sheet) into raw rows; quirks are reported, never hidden.
    analyse    a column type per column, checked against every value; clean
               names; a primary key per table; links between tables (kept only
               when every value matches).
    preview    what the browser shows and lets the user edit: table and column
               names, column types, date order, which tables to import.
    import     (see import_tables) one transaction into the workspace schema.

Owner decisions (2026-10-19): the preview is editable; ambiguous numeric
dates (03/04/2025) are read day first and flagged, switchable per column.
"""

import codecs
import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, Iterable, List, Optional, Tuple

MAX_FILES = 10
MAX_TOTAL_BYTES = 10 * 1024 * 1024
MAX_ROWS_PER_TABLE = 200_000
MAX_COLUMNS = 200
MAX_CELL_CHARS = 10_000
MAX_XLSX_UNCOMPRESSED = 100 * 1024 * 1024  # a 10 MB workbook that unpacks to more is a zip bomb
MAX_XLSX_RATIO = 200
TYPES = ("integer", "decimal", "date", "timestamp", "boolean", "text")

BIGINT_MIN, BIGINT_MAX = -(2**63), 2**63 - 1


class UploadError(Exception):
    """A file or plan problem the user can fix; the message says how."""


# --------------------------------------------------------------------------- #
# Names
# --------------------------------------------------------------------------- #

def clean_name(raw: Any, fallback: str) -> str:
    """'Order Date' -> 'order_date'; 'Été 2025' -> 'ete_2025'; '1st' -> 'c_1st'."""
    text = unicodedata.normalize("NFKD", str(raw or "")).encode("ascii", "ignore").decode()
    text = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", text.strip())  # camelCase -> camel_case
    text = re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").lower()
    if not text:
        text = fallback
    if text[0].isdigit():
        text = f"c_{text}"
    return text[:63]


def _dedupe(names: List[str]) -> List[str]:
    seen: Dict[str, int] = {}
    out = []
    for name in names:
        candidate, n = name, 1
        while candidate in seen:
            n += 1
            suffix = f"_{n}"
            candidate = f"{name[: 63 - len(suffix)]}{suffix}"
        seen[candidate] = 1
        out.append(candidate)
    return out


def singular(name: str) -> str:
    if name.endswith("ies") and len(name) > 3:
        return name[:-3] + "y"
    if name.endswith(("ses", "xes", "ches", "shes")):
        return name[:-2]
    if name.endswith("s") and not name.endswith("ss"):
        return name[:-1]
    return name


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #

@dataclass
class RawTable:
    name: str  # suggested table name
    source: str  # "orders.csv" or "book.xlsx › Sheet1"
    headers: List[str]
    rows: List[Tuple[Any, ...]]  # None for an empty cell; str, or a native Excel value
    notes: List[str] = field(default_factory=list)


def parse_files(files: List[Tuple[str, bytes]]) -> List[RawTable]:
    """Every table in the uploaded files, with names unique across all of them."""
    if not files:
        raise UploadError("Choose at least one CSV or Excel file.")
    if len(files) > MAX_FILES:
        raise UploadError(f"Upload at most {MAX_FILES} files at a time.")
    total = sum(len(data) for _, data in files)
    if total > MAX_TOTAL_BYTES:
        raise UploadError(f"The files add up to {total / 1048576:.1f} MB; the limit is {MAX_TOTAL_BYTES // 1048576} MB.")
    tables: List[RawTable] = []
    for filename, data in files:
        lower = filename.lower()
        if lower.endswith(".csv") or lower.endswith(".tsv") or lower.endswith(".txt"):
            tables.append(parse_csv(filename, data))
        elif lower.endswith(".xlsx") or lower.endswith(".xlsm"):
            tables.extend(parse_xlsx(filename, data))
        elif lower.endswith(".xls"):
            raise UploadError(f"{filename} is an old .xls workbook. Open it in Excel and save it as .xlsx.")
        elif lower.endswith(".sql"):
            raise UploadError(f"{filename}: SQL files can't be uploaded (they could run any statement). Export the tables as CSV instead.")
        else:
            raise UploadError(f"{filename}: only CSV and Excel (.xlsx) files can be uploaded.")
    names = _dedupe([t.name for t in tables])
    for table, name in zip(tables, names):
        table.name = name
    return tables


def _decode(data: bytes) -> Tuple[str, Optional[str]]:
    """Text of a CSV, and a note when it wasn't UTF-8."""
    if data.startswith(codecs.BOM_UTF16_LE) or data.startswith(codecs.BOM_UTF16_BE):
        return data.decode("utf-16"), "Read as UTF-16."
    try:
        return data.decode("utf-8-sig"), None
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1252"), "Not UTF-8: read as Windows-1252. Check accented characters in the preview."
    except UnicodeDecodeError:
        return data.decode("latin-1"), "Not UTF-8: read as Latin-1. Check accented characters in the preview."


def parse_csv(filename: str, data: bytes) -> RawTable:
    text, encoding_note = _decode(data)
    stem = filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1].rsplit(".", 1)[0]
    notes = [encoding_note] if encoding_note else []
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if filename.lower().endswith(".tsv") else ","
    if delimiter != ",":
        notes.append({"\t": "Columns are separated by tabs.", ";": "Columns are separated by semicolons.",
                      "|": "Columns are separated by | characters."}[delimiter])

    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    headers: Optional[List[str]] = None
    rows: List[Tuple[Any, ...]] = []
    short = skipped = 0
    first_skipped = None
    try:
        for line_no, cells in enumerate(reader, 1):
            if not any(c.strip() for c in cells):
                continue
            if headers is None:
                headers = [c.strip() for c in cells]
                continue
            values = [c.strip() or None for c in cells]
            if len(values) > len(headers):
                if any(v is not None for v in values[len(headers):]):
                    skipped += 1
                    first_skipped = first_skipped or line_no
                    continue
                values = values[: len(headers)]
            elif len(values) < len(headers):
                short += 1
                values += [None] * (len(headers) - len(values))
            rows.append(tuple(values))
            if len(rows) > MAX_ROWS_PER_TABLE:
                raise UploadError(f"{filename} has more than {MAX_ROWS_PER_TABLE:,} rows; split it or remove some.")
    except csv.Error as exc:
        raise UploadError(f"{filename} isn't a readable CSV file ({exc}).") from exc

    if headers is None:
        raise UploadError(f"{filename} is empty.")
    if short:
        notes.append(f"{short:,} row{'s' if short != 1 else ''} had fewer values than columns; the missing ones are empty.")
    if skipped:
        notes.append(
            f"{skipped:,} row{'s were' if skipped != 1 else ' was'} left out because "
            f"{'they have' if skipped != 1 else 'it has'} more values than there are columns "
            f"(first at line {first_skipped}); often an unquoted comma inside a value."
        )
    return _finish(clean_name(stem, "table"), filename, headers, rows, notes)


def _check_xlsx_container(filename: str, data: bytes) -> None:
    """Refuse zip bombs and XML with DOCTYPE/ENTITY (billion laughs, external entities)."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise UploadError(f"{filename} isn't a valid .xlsx file.") from exc
    with archive:
        infos = archive.infolist()
        unpacked = sum(i.file_size for i in infos)
        if unpacked > MAX_XLSX_UNCOMPRESSED or unpacked > MAX_XLSX_RATIO * max(len(data), 1):
            raise UploadError(f"{filename} unpacks to {unpacked / 1048576:.0f} MB, which is too large to read safely.")
        for info in infos:
            if info.filename.endswith((".xml", ".rels", ".vml")):
                head = archive.read(info.filename)
                if b"<!DOCTYPE" in head or b"<!ENTITY" in head:
                    raise UploadError(f"{filename} contains XML declarations Excel never writes, so it wasn't opened.")


def parse_xlsx(filename: str, data: bytes) -> List[RawTable]:
    import openpyxl  # only needed for Excel uploads

    _check_xlsx_container(filename, data)
    stem = clean_name(filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1].rsplit(".", 1)[0], "workbook")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise UploadError(f"{filename} couldn't be opened as an Excel workbook.") from exc
    tables = []
    try:
        sheets = workbook.worksheets
        for sheet in sheets:
            headers: Optional[List[str]] = None
            rows: List[Tuple[Any, ...]] = []
            for cells in sheet.iter_rows(values_only=True):
                if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
                    continue
                if headers is None:
                    headers = ["" if c is None else str(c).strip() for c in cells]
                    while headers and not headers[-1]:
                        headers.pop()  # trailing empty header cells
                    continue
                values = [(c.strip() or None) if isinstance(c, str) else c for c in cells[: len(headers)]]
                values += [None] * (len(headers) - len(values))
                rows.append(tuple(values))
                if len(rows) > MAX_ROWS_PER_TABLE:
                    raise UploadError(f"{filename} › {sheet.title} has more than {MAX_ROWS_PER_TABLE:,} rows.")
            if headers is None or not headers:
                continue  # an empty sheet
            notes = []
            name = stem if len(sheets) == 1 else clean_name(sheet.title, "sheet")
            tables.append(_finish(name, f"{filename} › {sheet.title}", headers, rows, notes))
    finally:
        workbook.close()
    if not tables:
        raise UploadError(f"{filename} has no data in any sheet.")
    return tables


def _finish(name: str, source: str, headers: List[str], rows, notes) -> RawTable:
    if len(headers) > MAX_COLUMNS:
        raise UploadError(f"{source} has {len(headers)} columns; the limit is {MAX_COLUMNS}.")
    for row in rows:
        for value in row:
            if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
                raise UploadError(f"{source} has a value longer than {MAX_CELL_CHARS:,} characters.")
    if not rows:
        notes.append("The file has a header row but no data.")
    return RawTable(name=name, source=source, headers=headers, rows=rows, notes=notes)


# --------------------------------------------------------------------------- #
# Values and types
# --------------------------------------------------------------------------- #

_INT = re.compile(r"^[+-]?(0|[1-9]\d*)$")
_DEC = re.compile(r"^[+-]?(0|[1-9]\d*)?(\.\d+)?([eE][+-]?\d+)?$")
_GROUPED = re.compile(r"^[+-]?[1-9]\d{0,2}(,\d{3})+(\.\d+)?$")  # 1,234,567.89
_BOOL = {"true": True, "false": False, "yes": True, "no": False, "t": True, "f": False, "y": True, "n": False}
_ISO_DATE = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$")
_NUM_DATE = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})$")
_MONTHS = {
    name: number
    for number, names in enumerate([
        ("jan", "january"), ("feb", "february"), ("mar", "march"), ("apr", "april"), ("may",), ("jun", "june"),
        ("jul", "july"), ("aug", "august"), ("sep", "sept", "september"), ("oct", "october"),
        ("nov", "november"), ("dec", "december"),
    ], 1)
    for name in names
}
_NAMED_DATE = re.compile(r"^(?:(\d{1,2})[\s\-]+([A-Za-z]{3,9})\.?[\s\-,]+(\d{4})|([A-Za-z]{3,9})\.?\s+(\d{1,2}),?\s+(\d{4}))$")
_TIME = r"(\d{1,2}):(\d{2})(?::(\d{2})(?:\.(\d{1,6}))?)?\s*(Z|[+-]\d{2}:?\d{2})?"
_ISO_TS = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})[T ]" + _TIME + r"$")
_NUM_TS = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})[ T]" + _TIME + r"$")


def _year(text: str) -> int:
    y = int(text)
    return y if len(text) == 4 else (2000 + y if y < 70 else 1900 + y)


def _month_name(text: str) -> Optional[int]:
    """'Apr', 'April' or 'Sept' -> its month number; anything else -> None."""
    return _MONTHS.get(text.lower())


def _safe_date(y: int, m: int, d: int) -> Optional[date]:
    try:
        return date(y, m, d)
    except ValueError:
        return None


def parse_date(text: str, order: str = "dmy") -> Optional[date]:
    """A date from text, reading ambiguous numeric dates in ``order`` ('dmy' or 'mdy')."""
    m = _ISO_DATE.match(text)
    if m:
        return _safe_date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = _NUM_DATE.match(text)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), _year(m.group(3))
        day, month = (a, b) if order == "dmy" else (b, a)
        return _safe_date(y, month, day)
    m = _NAMED_DATE.match(text)
    if m:
        if m.group(1):
            day, month, y = int(m.group(1)), _month_name(m.group(2)), int(m.group(3))
        else:
            month, day, y = _month_name(m.group(4)), int(m.group(5)), int(m.group(6))
        return _safe_date(y, month, day) if month else None
    return None


def parse_timestamp(text: str, order: str = "dmy") -> Optional[datetime]:
    m = _ISO_TS.match(text)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        rest = m.groups()[3:]
    else:
        m = _NUM_TS.match(text)
        if not m:
            return None
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        d, mo = (a, b) if order == "dmy" else (b, a)
        rest = m.groups()[3:]
    hour, minute, second, fraction, zone = rest
    try:
        value = datetime(y, mo, d, int(hour), int(minute), int(second or 0), int((fraction or "0").ljust(6, "0")))
    except ValueError:
        return None
    if zone:
        if zone == "Z":
            offset = 0
        else:
            sign = 1 if zone[0] == "+" else -1
            digits = zone[1:].replace(":", "")
            offset = sign * (int(digits[:2]) * 60 + int(digits[2:]))
        value -= timedelta(minutes=offset)  # stored as UTC
    return value


def _date_parts(text: str) -> Optional[Tuple[int, int]]:
    """For a numeric d/m/y-style date: its first two numbers, to work out the order."""
    m = _NUM_DATE.match(text) or _NUM_TS.match(text)
    return (int(m.group(1)), int(m.group(2))) if m else None


@dataclass
class ColumnAnalysis:
    type: str
    nullable: bool
    date_order: Optional[str] = None  # 'dmy' / 'mdy' for numeric dates
    ambiguous_dates: bool = False
    notes: List[str] = field(default_factory=list)


def analyse_column(values: Iterable[Any]) -> ColumnAnalysis:
    """The narrowest type every non-empty value fits."""
    possible = {"boolean", "integer", "decimal", "date", "timestamp"}
    nullable = False
    any_value = grouped = False
    day_first_proof = month_first_proof = False
    numeric_dates = False
    has_time = has_zone = False
    for value in values:
        if value is None:
            nullable = True
            continue
        any_value = True
        if isinstance(value, bool):
            possible &= {"boolean"}
        elif isinstance(value, int):
            possible &= {"integer", "decimal"} if BIGINT_MIN <= value <= BIGINT_MAX else {"decimal"}
        elif isinstance(value, float):
            possible &= {"integer", "decimal"} if value.is_integer() and abs(value) < 2**53 else {"decimal"}
        elif isinstance(value, datetime):
            possible &= {"timestamp", "date"} if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0) else {"timestamp"}
            has_time = has_time or value.time() != datetime.min.time()
        elif isinstance(value, date):
            possible &= {"date", "timestamp"}
        else:
            text = str(value)
            keep = set()
            if "boolean" in possible and text.lower() in _BOOL:
                keep.add("boolean")
            if _INT.match(text):
                if BIGINT_MIN <= int(text) <= BIGINT_MAX:
                    keep |= {"integer", "decimal"}
            elif _DEC.match(text) and any(ch.isdigit() for ch in text) and not re.match(r"^[+-]?0\d", text):
                keep.add("decimal")
            elif _GROUPED.match(text):
                keep.add("decimal")
                grouped = True
            if possible & {"date", "timestamp"}:
                parts = _date_parts(text)
                if parts:
                    numeric_dates = True
                    a, b = parts
                    day_first_proof |= a > 12
                    month_first_proof |= b > 12
                if parse_date(text, "dmy") or parse_date(text, "mdy"):
                    keep |= {"date", "timestamp"}
                elif parse_timestamp(text, "dmy") or parse_timestamp(text, "mdy"):
                    keep.add("timestamp")
                    has_time = True
                    has_zone = has_zone or bool(re.search(r"(Z|[+-]\d{2}:?\d{2})$", text))
            possible &= keep
        if not possible:
            break

    if not any_value:
        return ColumnAnalysis("text", True, notes=["Every value is empty."])
    order = None
    ambiguous = False
    if numeric_dates and possible & {"date", "timestamp"}:
        if day_first_proof and month_first_proof:
            possible -= {"date", "timestamp"}  # 25/12 and 12/25 in one column: not a date column
        elif month_first_proof:
            order = "mdy"
        else:
            order, ambiguous = "dmy", not day_first_proof
    for kind in ("boolean", "integer", "decimal", "date", "timestamp"):
        if kind in possible and (kind != "date" or not has_time):
            notes = []
            if kind == "decimal" and grouped:
                notes.append("Numbers with thousands separators (1,234.5) were read as numbers.")
            if kind == "timestamp" and has_zone:
                notes.append("Times with a timezone were converted to UTC.")
            if kind in ("date", "timestamp") and ambiguous:
                notes.append("Dates like 03/04/2025 could be day-first or month-first; read as day first (DD/MM).")
            return ColumnAnalysis(kind, nullable, order if kind in ("date", "timestamp") else None,
                                  ambiguous if kind in ("date", "timestamp") else False, notes)
    return ColumnAnalysis("text", nullable)


def convert(value: Any, kind: str, date_order: Optional[str] = "dmy") -> Any:
    """``value`` as a Python value for column type ``kind``; raises ValueError if it doesn't fit."""
    if value is None:
        return None
    order = date_order or "dmy"
    if kind == "text":
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, datetime):
            return value.isoformat(sep=" ")
        return str(value)
    if kind == "boolean":
        if isinstance(value, bool):
            return value
        if isinstance(value, str) and value.lower() in _BOOL:
            return _BOOL[value.lower()]
        raise ValueError
    if kind == "integer":
        if isinstance(value, bool):
            raise ValueError
        if isinstance(value, int):
            result = value
        elif isinstance(value, float) and value.is_integer():
            result = int(value)
        elif isinstance(value, str) and _INT.match(value):
            result = int(value)
        else:
            raise ValueError
        if not BIGINT_MIN <= result <= BIGINT_MAX:
            raise ValueError
        return result
    if kind == "decimal":
        if isinstance(value, bool):
            raise ValueError
        if isinstance(value, (int, float)):
            return Decimal(str(value))
        text = str(value)
        if _GROUPED.match(text):
            text = text.replace(",", "")
        elif not (_DEC.match(text) and any(ch.isdigit() for ch in text)):
            raise ValueError
        try:
            return Decimal(text)
        except InvalidOperation as exc:
            raise ValueError from exc
    if kind == "date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        parsed = parse_date(str(value), order)
        if parsed is None:
            raise ValueError
        return parsed
    if kind == "timestamp":
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime(value.year, value.month, value.day)
        text = str(value)
        parsed = parse_timestamp(text, order)
        if parsed is None:
            day = parse_date(text, order)
            if day is None:
                raise ValueError
            parsed = datetime(day.year, day.month, day.day)
        return parsed
    raise ValueError(f"unknown type {kind}")


# --------------------------------------------------------------------------- #
# Analysis -> plan (what the preview shows and the user edits)
# --------------------------------------------------------------------------- #

@dataclass
class ColumnPlan:
    key: str
    original: str
    name: str
    type: str
    nullable: bool
    date_order: Optional[str]
    ambiguous_dates: bool
    samples: List[str]
    notes: List[str]


@dataclass
class TablePlan:
    key: str
    name: str
    source: str
    include: bool
    row_count: int
    columns: List[ColumnPlan]
    primary_key: Optional[str]
    notes: List[str]


@dataclass
class Link:
    table: str
    column: str
    ref_table: str
    ref_column: str


@dataclass
class Plan:
    tables: List[TablePlan]
    links: List[Link]
    notes: List[str]


def _samples(values: List[Any], limit: int = 4) -> List[str]:
    out: List[str] = []
    for value in values:
        if value is None:
            continue
        text = convert(value, "text")
        if text not in out:
            out.append(text if len(text) <= 40 else text[:39] + "…")
        if len(out) == limit:
            break
    return out


def analyse(tables: List[RawTable]) -> Plan:
    plans = []
    for t_index, raw in enumerate(tables):
        names = _dedupe([clean_name(h, f"column_{i + 1}") for i, h in enumerate(raw.headers)])
        columns = []
        for c_index, (original, name) in enumerate(zip(raw.headers, names)):
            values = [row[c_index] for row in raw.rows]
            info = analyse_column(values)
            columns.append(ColumnPlan(
                key=f"c{c_index}", original=original or f"(column {c_index + 1})", name=name, type=info.type,
                nullable=info.nullable, date_order=info.date_order, ambiguous_dates=info.ambiguous_dates,
                samples=_samples(values), notes=info.notes,
            ))
        notes = list(raw.notes)
        if any(not c.original.strip() or c.original.startswith("(column") for c in columns):
            notes.append("Some columns had no header and were named column_N.")
        plans.append(TablePlan(key=f"t{t_index}", name=raw.name, source=raw.source, include=True,
                               row_count=len(raw.rows), columns=columns, primary_key=None, notes=notes))
    plan = Plan(tables=plans, links=[], notes=[])
    finalize(plan, tables)
    return plan


def _column_values(raw: RawTable, column: ColumnPlan) -> List[Any]:
    index = int(column.key[1:])
    return [row[index] for row in raw.rows]


def finalize(plan: Plan, raws: List[RawTable]) -> None:
    """(Re)compute primary keys and links for the plan as it is now (after any edits)."""
    by_key = {raw_key: raw for raw_key, raw in zip((f"t{i}" for i in range(len(raws))), raws)}
    included = [t for t in plan.tables if t.include]
    keys: Dict[str, Tuple[ColumnPlan, set]] = {}
    for table in included:
        table.primary_key = None
        raw = by_key[table.key]
        candidates = [c for c in table.columns if c.type in ("integer", "text")]
        # The key is named after what a row is: "id", "<table>_id", or after the
        # file it came from (renaming employees to staff keeps employee_id), or
        # a first column ending in _id (as exports usually have it).
        preferred = ["id"]
        for stem in dict.fromkeys([table.name, raw.name]):
            preferred += [f"{stem}_id", f"{singular(stem)}_id", f"{singular(stem)}id"]
        first = table.columns[0].name if table.columns else ""
        if first.endswith("_id") and first not in preferred:
            preferred.append(first)
        candidates.sort(key=lambda c: preferred.index(c.name) if c.name in preferred else len(preferred))
        for column in candidates:
            if column.name not in preferred:
                break
            values = [convert(v, column.type, column.date_order) for v in _column_values(raw, column)]
            if values and None not in values and len(set(values)) == len(values):
                table.primary_key = column.name
                keys[table.name] = (column, set(values))
                break

    plan.links = []
    plan.notes = [n for n in plan.notes if not n.startswith("Not linked:")]
    for table in included:
        raw = by_key[table.key]
        for column in table.columns:
            if column.name == table.primary_key:
                continue
            for target in included:
                if target is table or target.name not in keys:
                    continue
                pk, pk_values = keys[target.name]
                stem = singular(target.name)
                names = {f"{stem}_id", f"{stem}id", f"{target.name}_id"}
                if pk.name != "id":
                    names.add(pk.name)  # orders.customer_id -> customers.customer_id; a bare "id" never links
                if column.name not in names:
                    continue
                if column.type != pk.type:
                    continue
                values = [convert(v, column.type, column.date_order) for v in _column_values(raw, column)]
                unmatched = sum(1 for v in values if v is not None and v not in pk_values)
                if unmatched:
                    plan.notes.append(
                        f"Not linked: {table.name}.{column.name} looks like a link to {target.name}.{pk.name}, "
                        f"but {unmatched:,} value{'s have' if unmatched != 1 else ' has'} no match there."
                    )
                    continue
                plan.links.append(Link(table.name, column.name, target.name, pk.name))
                break


def to_preview(plan: Plan) -> dict:
    return {
        "tables": [
            {
                "key": t.key, "name": t.name, "source": t.source, "include": t.include, "rows": t.row_count,
                "primary_key": t.primary_key, "notes": t.notes,
                "columns": [
                    {
                        "key": c.key, "original": c.original, "name": c.name, "type": c.type, "nullable": c.nullable,
                        "date_order": c.date_order, "ambiguous_dates": c.ambiguous_dates, "samples": c.samples,
                        "notes": c.notes,
                    }
                    for c in t.columns
                ],
            }
            for t in plan.tables
        ],
        "links": [{"table": l.table, "column": l.column, "ref_table": l.ref_table, "ref_column": l.ref_column} for l in plan.links],
        "notes": plan.notes,
        "types": list(TYPES),
        "limits": {"files": MAX_FILES, "total_mb": MAX_TOTAL_BYTES // 1048576, "rows_per_table": MAX_ROWS_PER_TABLE},
    }


def apply_edits(plan: Plan, raws: List[RawTable], edits: Optional[dict]) -> None:
    """
    Apply the user's preview edits (by table/column key), check every value
    fits its (possibly changed) type, then recompute keys and links. Raises
    UploadError listing every problem, so the user can fix them in one go.
    """
    problems: List[str] = []
    if edits:
        tables = {t.key: t for t in plan.tables}
        for edit in edits.get("tables") or []:
            table = tables.get(str(edit.get("key")))
            if table is None:
                continue
            if "include" in edit:
                table.include = bool(edit["include"])
            if edit.get("name"):
                table.name = clean_name(edit["name"], table.name)
            columns = {c.key: c for c in table.columns}
            for cedit in edit.get("columns") or []:
                column = columns.get(str(cedit.get("key")))
                if column is None:
                    continue
                if cedit.get("name"):
                    column.name = clean_name(cedit["name"], column.name)
                if cedit.get("type") in TYPES:
                    column.type = cedit["type"]
                if cedit.get("date_order") in ("dmy", "mdy"):
                    column.date_order = cedit["date_order"]
    included = [t for t in plan.tables if t.include]
    if not included:
        raise UploadError("Choose at least one table to import.")
    table_names = [t.name for t in included]
    for name in {n for n in table_names if table_names.count(n) > 1}:
        problems.append(f"Two tables are named {name}; rename one.")
    raws_by_key = {f"t{i}": raw for i, raw in enumerate(raws)}
    for table in included:
        names = [c.name for c in table.columns]
        for name in {n for n in names if names.count(n) > 1}:
            problems.append(f"{table.name}: two columns are named {name}; rename one.")
        raw = raws_by_key[table.key]
        for column in table.columns:
            bad, example = 0, None
            for value in _column_values(raw, column):
                try:
                    convert(value, column.type, column.date_order)
                except ValueError:
                    bad += 1
                    example = example if example is not None else value
            if bad:
                problems.append(
                    f"{table.name}.{column.name}: {bad:,} value{'s don' if bad != 1 else ' doesn'}'t fit "
                    f"the type {column.type} (for example {str(example)[:40]!r}). Choose another type."
                )
            if column.type in ("date", "timestamp") and not column.date_order:
                column.date_order = "dmy"
    if problems:
        raise UploadError(" ".join(problems))
    finalize(plan, raws)


# --------------------------------------------------------------------------- #
# Import (inside the workspace's creation transaction; see WorkspaceStore.create)
# --------------------------------------------------------------------------- #

PG_TYPES = {"integer": "BIGINT", "decimal": "NUMERIC", "date": "DATE", "timestamp": "TIMESTAMP",
            "boolean": "BOOLEAN", "text": "TEXT"}


def _megabytes(env: str, default: int) -> int:
    import os
    try:
        value = int(os.getenv(env, default))
    except ValueError:
        return default * 1048576
    return (value if value > 0 else default) * 1048576


def user_quota_bytes() -> int:
    """Database storage one user's uploaded data may take (env UPLOAD_QUOTA_MB, default 25)."""
    return _megabytes("UPLOAD_QUOTA_MB", 25)


def total_quota_bytes() -> int:
    """All workspaces together (env UPLOAD_TOTAL_MB, default 300), so uploads never crowd out the demo."""
    return _megabytes("UPLOAD_TOTAL_MB", 300)


def import_tables(conn, schema: str, plan: Plan, raws: List[RawTable]) -> dict:
    """
    Create the plan's tables in ``schema``, load every row, add the links and
    check the storage limits. Raises UploadError (the caller's transaction
    then rolls everything back).
    """
    from psycopg import sql

    raws_by_key = {f"t{i}": raw for i, raw in enumerate(raws)}
    summary = []
    for table in [t for t in plan.tables if t.include]:
        columns = [
            sql.SQL("{} {}").format(sql.Identifier(c.name), sql.SQL(PG_TYPES[c.type])) for c in table.columns
        ]
        if table.primary_key:
            columns.append(sql.SQL("PRIMARY KEY ({})").format(sql.Identifier(table.primary_key)))
        conn.execute(sql.SQL("CREATE TABLE {}.{} ({})").format(
            sql.Identifier(schema), sql.Identifier(table.name), sql.SQL(", ").join(columns)))
        raw = raws_by_key[table.key]
        indexes = [int(c.key[1:]) for c in table.columns]
        copy_sql = sql.SQL("COPY {}.{} ({}) FROM STDIN").format(
            sql.Identifier(schema), sql.Identifier(table.name),
            sql.SQL(", ").join(sql.Identifier(c.name) for c in table.columns))
        with conn.cursor() as cur, cur.copy(copy_sql) as copy:
            for row in raw.rows:
                copy.write_row([convert(row[i], c.type, c.date_order) for i, c in zip(indexes, table.columns)])
        summary.append({"name": table.name, "rows": len(raw.rows), "columns": len(table.columns),
                        "primary_key": table.primary_key})
    for link in plan.links:
        conn.execute(sql.SQL("ALTER TABLE {}.{} ADD FOREIGN KEY ({}) REFERENCES {}.{} ({})").format(
            sql.Identifier(schema), sql.Identifier(link.table), sql.Identifier(link.column),
            sql.Identifier(schema), sql.Identifier(link.ref_table), sql.Identifier(link.ref_column)))
    for table in summary:
        conn.execute(sql.SQL("ANALYZE {}.{}").format(sql.Identifier(schema), sql.Identifier(table["name"])))

    size = conn.execute(
        """
        SELECT COALESCE(sum(pg_total_relation_size(c.oid)), 0) AS bytes
        FROM pg_class c JOIN pg_namespace n ON n.oid = c.relnamespace
        WHERE n.nspname = %s AND c.relkind = 'r'
        """,
        (schema,),
    ).fetchone()["bytes"]
    if size > user_quota_bytes():
        raise UploadError(
            f"Your data takes {size / 1048576:.1f} MB in the database; the limit is "
            f"{user_quota_bytes() // 1048576} MB. Upload fewer rows or columns."
        )
    everyone = conn.execute(
        r"""
        SELECT COALESCE(sum(pg_total_relation_size(c.oid)), 0) AS bytes
        FROM pg_class c JOIN pg_namespace n ON n.oid = c.relnamespace
        WHERE n.nspname LIKE 'ws\_%' AND c.relkind = 'r'
        """
    ).fetchone()["bytes"]
    if everyone > total_quota_bytes():
        raise UploadError("Storage for uploaded data is full right now. Try again later, or upload less.")
    return {"tables": summary, "links": [vars(l) for l in plan.links], "size_bytes": int(size)}
